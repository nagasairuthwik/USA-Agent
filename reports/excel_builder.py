"""
Excel Builder Module
Generates the corporate styled 6-sheet OpenPyXL master workbook.
Preserves exact technical specification headers, navy executive theme, auto-filters,
safe string formatting (formula injection prevention), and atomic write safety.
"""

import os
import tempfile
import shutil
from typing import List, Dict, Any, Optional
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from core.models import (
    CityMaster,
    TransitConnectivity,
    Accommodation,
    SafetyIndex,
    RecommendationScore,
    FinalOutput
)

class ExcelWorkbookBuilder:
    def __init__(self, budget_min: float = 50.0, budget_max: float = 60.0, max_commute_minutes: int = 90):
        self.budget_min = float(budget_min)
        self.budget_max = float(budget_max)
        self.max_commute_minutes = int(max_commute_minutes)

    def _sanitize_cell_value(self, val: Any) -> Any:
        """Prevents formula injection in Excel by prefixing dangerous formula symbols."""
        if isinstance(val, str) and len(val) > 0:
            if val[0] in ("=", "+", "-", "@"):
                return "'" + val
        return val

    def build_workbook(
        self,
        filepath: str,
        cities: List[CityMaster],
        transits: List[TransitConnectivity],
        accomms: List[Accommodation],
        safeties: List[SafetyIndex],
        rec_scores: List[RecommendationScore],
        final_outputs: List[FinalOutput]
    ):
        """
        Builds the complete 6-sheet workbook and saves it atomically.
        """
        wb = openpyxl.Workbook()
        wb.remove(wb.active)

        TITLE_FONT = Font(name="Segoe UI", size=14, bold=True, color="FFFFFF")
        TITLE_FILL = PatternFill(start_color="0F172A", end_color="0F172A", fill_type="solid")
        THEME_HEADER_FONT = Font(name="Segoe UI", size=10, bold=True, color="FFFFFF")
        THEME_HEADER_FILL = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
        REGULAR_FONT = Font(name="Segoe UI", size=9, color="1E293B")
        ZEBRA_FILL = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")
        HIGHLIGHT_TOP1_FILL = PatternFill(start_color="ECFDF5", end_color="ECFDF5", fill_type="solid")
        THIN_BORDER = Border(
            left=Side(style="thin", color="E2E8F0"),
            right=Side(style="thin", color="E2E8F0"),
            top=Side(style="thin", color="E2E8F0"),
            bottom=Side(style="thin", color="E2E8F0")
        )

        def _write_table(ws, title: str, headers: List[str], data: List[List[Any]], number_formats: Optional[Dict[int, str]] = None):
            ws.views.sheetView[0].showGridLines = True
            num_cols = len(headers)

            ws.row_dimensions[1].height = 36
            ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=num_cols)
            title_cell = ws.cell(row=1, column=1, value=title)
            title_cell.font = TITLE_FONT
            title_cell.fill = TITLE_FILL
            title_cell.alignment = Alignment(horizontal="center", vertical="center")

            ws.row_dimensions[2].height = 8

            header_row = 3
            ws.row_dimensions[header_row].height = 26
            for col_idx, h in enumerate(headers, start=1):
                cell = ws.cell(row=header_row, column=col_idx, value=h)
                cell.font = THEME_HEADER_FONT
                cell.fill = THEME_HEADER_FILL
                cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
                cell.border = THIN_BORDER

            start_data_row = 4
            if len(data) == 0:
                ws.row_dimensions[start_data_row].height = 28
                ws.merge_cells(start_row=start_data_row, start_column=1, end_row=start_data_row, end_column=num_cols)
                msg = "No hotels found within your selected budget"
                empty_cell = ws.cell(row=start_data_row, column=1, value=msg)
                empty_cell.font = Font(name="Segoe UI", size=10, italic=True, color="64748B")
                empty_cell.alignment = Alignment(horizontal="center", vertical="center")
                last_row = start_data_row
            else:
                for row_offset, row_data in enumerate(data):
                    curr_row = start_data_row + row_offset
                    ws.row_dimensions[curr_row].height = 21
                    is_zebra = (row_offset % 2 == 1)
                    is_rank1 = False
                    if "Rank" in headers and row_data[headers.index("Rank")] == 1:
                        is_rank1 = True
                    elif "Recommendation" in headers and "⭐ #1" in str(row_data[-1]):
                        is_rank1 = True

                    for col_idx, raw_val in enumerate(row_data, start=1):
                        val = self._sanitize_cell_value(raw_val)
                        cell = ws.cell(row=curr_row, column=col_idx, value=val)
                        cell.font = REGULAR_FONT
                        cell.border = THIN_BORDER
                        if is_rank1:
                            cell.fill = HIGHLIGHT_TOP1_FILL
                        elif is_zebra:
                            cell.fill = ZEBRA_FILL

                        if isinstance(val, (int, float)):
                            cell.alignment = Alignment(horizontal="right", vertical="center")
                        else:
                            cell.alignment = Alignment(horizontal="left", vertical="center")
                        if number_formats and col_idx in number_formats:
                            cell.number_format = number_formats[col_idx]

                last_row = start_data_row + len(data) - 1

            ws.auto_filter.ref = f"A{header_row}:{get_column_letter(num_cols)}{last_row}"
            for col in ws.columns:
                col_letter = get_column_letter(col[0].column)
                max_len = max([len(str(c.value or '')) for c in col if c.row >= header_row] + [10])
                ws.column_dimensions[col_letter].width = min(max(max_len + 3, 12), 65)

        # Sheet 1: City_Master
        ws1 = wb.create_sheet(title="City_Master")
        _write_table(
            ws1,
            "Sheet 1 - City Master Reference Data",
            ["City_ID", "City", "State", "Downtown_District", "Primary_Transit_System", "Latitude", "Longitude", "Google_Maps_District_URL", "Last_Updated"],
            [[c.city_id, c.city, c.state, c.downtown_district, c.primary_transit_system, c.latitude, c.longitude, c.google_maps_district_url, c.last_updated] for c in cities],
            {6: "0.0000", 7: "0.0000"}
        )

        # Sheet 2: Transit_Connectivity
        ws2 = wb.create_sheet(title="Transit_Connectivity")
        _write_table(
            ws2,
            f"Sheet 2 - Commuter Rail & Metro Transit Connectivity (Commute <= {self.max_commute_minutes}m)",
            ["City", "Downtown", "Suburb", "Transit_System", "Transit_Mode", "Nearest_Station", "Commute_Minutes", "Transfers", "Fare_USD", "Google_Maps_Station_URL"],
            [[t.city, t.downtown, t.suburb, t.transit_system, t.transit_mode, t.nearest_station, t.commute_minutes, t.transfers, t.fare_usd, t.google_maps_station_url] for t in transits],
            {7: "0", 8: "0", 9: "$#,##0.00"}
        )

        # Sheet 3: Accommodation_Master
        ws3 = wb.create_sheet(title="Accommodation_Master")
        _write_table(
            ws3,
            f"Sheet 3 - Verified Accommodations (Target Tier: USD ${self.budget_min:.0f} - ${self.budget_max:.0f} / Night)",
            [
                "City", "Locality", "Hotel_Name", "Address", "Currency",
                "Base_Nightly_Rate_USD", "Taxes_And_Fees_USD", "Total_Booking_USD",
                "Availability_Status", "Check_In_Date", "Check_Out_Date", "Occupancy",
                "Google_Hotels_URL", "Google_Maps_URL", "Booking_Demand_API_URL", "Expedia_Rapid_API_URL", "Wyndham_Official_URL",
                "Primary_Price_Source", "Retrieval_Timestamp",
                "Rating", "Review_Count", "Distance_To_Station_Miles"
            ],
            [
                [
                    a.city, a.locality, a.hotel_name, a.address, a.currency,
                    a.base_nightly_rate_usd, a.taxes_and_fees_usd, a.total_booking_usd,
                    a.availability_status, a.check_in_date, a.check_out_date, a.occupancy,
                    a.google_hotels_url, a.google_maps_url, a.booking_demand_api_url, a.expedia_rapid_api_url, a.wyndham_official_url,
                    a.price_source, a.retrieval_timestamp,
                    a.rating, a.review_count, a.distance_to_station_miles
                ]
                for a in accomms
            ],
            {6: "$#,##0.00", 7: "$#,##0.00", 8: "$#,##0.00", 20: "0.0", 21: "#,##0", 22: "0.00"}
        )

        # Sheet 4: Safety_Index
        ws4 = wb.create_sheet(title="Safety_Index")
        _write_table(
            ws4,
            "Sheet 4 - Neighborhood Safety & Crime Intelligence Index",
            ["City", "Locality", "Safety_Score", "Crime_Index", "Walkability", "Category"],
            [[s.city, s.locality, s.safety_score, s.crime_index, s.walkability, s.category] for s in safeties],
            {3: "0.0", 4: "0.0", 5: "0"}
        )

        # Sheet 5: Recommendation_Score
        ws5 = wb.create_sheet(title="Recommendation_Score")
        _write_table(
            ws5,
            "Sheet 5 - Multi-Factor Recommendation & Ranking Scores",
            ["City", "Locality", "Safety_Score", "Cost_Score", "Commute_Score", "Transit_Score", "Overall_Score", "Rank"],
            [[r.city, r.locality, r.safety_score, r.cost_score, r.commute_score, r.transit_score, r.overall_score, r.rank] for r in rec_scores],
            {3: "0.0", 4: "0.0", 5: "0.0", 6: "0.0", 7: "0.00", 8: "0"}
        )

        # Sheet 6: Final_Output
        ws6 = wb.create_sheet(title="Final_Output")
        _write_table(
            ws6,
            f"Sheet 6 - Executive Travel Recommendations (${self.budget_min:.0f}-${self.budget_max:.0f} Budget Tier & Direct Rail)",
            [
                "City", "Downtown", "Recommended_Locality", "Transit_System", "Station",
                "Commute_Time", "Hotel", "Currency", "Base_Nightly_Rate", "Taxes_And_Fees",
                "Total_Booking_Cost", "Availability",
                "Google_Hotels_Comparison", "Google_Maps_Location", "Booking_Demand_API", "Expedia_Rapid_API", "Wyndham_Official_Link",
                "Safety_Score", "Overall_Score",
                "Primary_Price_Source", "Retrieval_Timestamp", "Recommendation"
            ],
            [
                [
                    f.city, f.downtown, f.recommended_locality, f.transit_system, f.station,
                    f.commute_time, f.hotel, f.currency, f.rate, f.taxes_and_fees,
                    f.total_booking, f.availability_status,
                    f.google_hotels_url, f.google_maps_url, f.booking_demand_api_url, f.expedia_rapid_api_url, f.wyndham_official_url,
                    f.safety, f.overall_score,
                    f.price_source, f.retrieval_timestamp, f.recommendation
                ]
                for f in final_outputs
            ],
            {9: "$#,##0.00", 10: "$#,##0.00", 11: "$#,##0.00", 18: "0.0", 19: "0.00"}
        )

        # Atomic file write: save to temporary file first, then atomically replace
        target_dir = os.path.dirname(os.path.abspath(filepath))
        os.makedirs(target_dir, exist_ok=True)
        temp_fd, temp_path = tempfile.mkstemp(suffix=".xlsx", dir=target_dir)
        os.close(temp_fd)
        try:
            wb.save(temp_path)
            shutil.move(temp_path, filepath)
        except Exception:
            if os.path.exists(temp_path):
                os.remove(temp_path)
            raise
